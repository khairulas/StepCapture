"""Export an edited session to Word (.docx), Markdown and HTML."""
from __future__ import annotations

import base64
import html
import io
import re
from datetime import datetime
from pathlib import Path

from .project import Project
from .render import render_step

FORMATS = ("docx", "md", "html")
_INLINE = re.compile(r"(\*\*[^*]+?\*\*|\*[^*\s][^*]*?\*)")


def slug(text: str) -> str:
    s = re.sub(r"[^\w\- ]+", "", text).strip().replace(" ", "_")
    return s[:60] or "tutorial"


def inline_parts(text: str):
    """Split 'Click **Save** in *Word*' into [(text, bold, italic), ...]."""
    parts = []
    for tok in _INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            parts.append((tok[2:-2], True, False))
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            parts.append((tok[1:-1], False, True))
        else:
            parts.append((tok, False, False))
    return parts


def _date_line(project: Project) -> str:
    d = project.recorded_on or datetime.now()
    return d.strftime("%d %B %Y")


def _rendered(project: Project):
    """Yield (number, step, image-or-None) in tutorial order."""
    for n, step in enumerate(project.steps, start=1):
        yield n, step, render_step(project.image_path(step), step, border=True)


def export_dir(project: Project) -> Path:
    out = project.folder / "export"
    out.mkdir(exist_ok=True)
    return out


# --------------------------------------------------------------------- Markdown
def export_markdown(project: Project) -> Path:
    out = export_dir(project)
    images = out / "images"
    images.mkdir(exist_ok=True)
    name = slug(project.display_title())

    lines = [f"# {project.display_title()}", "", f"*{_date_line(project)}*", ""]
    if project.intro:
        lines += [project.intro, ""]
    for n, step, img in _rendered(project):
        lines += [f"## Step {n}", "", Project.instruction(step), ""]
        if img is not None:
            fname = f"step_{n:03d}.png"
            img.save(images / fname)
            lines += [f"![Step {n}](images/{fname})", ""]
    path = out / f"{name}.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


# ------------------------------------------------------------------------- HTML
_CSS = """
:root { --fg:#1f2328; --muted:#59636e; --accent:#c62828; --line:#d8dee4; --bg:#ffffff; }
* { box-sizing:border-box; }
body { margin:0; background:var(--bg); color:var(--fg);
       font:16px/1.6 "Segoe UI", system-ui, -apple-system, Roboto, Arial, sans-serif; }
main { max-width:960px; margin:0 auto; padding:40px 20px 80px; }
h1 { font-size:2rem; line-height:1.2; margin:0 0 4px; }
.date { color:var(--muted); margin:0 0 24px; }
.intro { margin:0 0 32px; white-space:pre-line; }
.step { border-top:1px solid var(--line); padding:24px 0; break-inside:avoid; page-break-inside:avoid; }
.step h2 { display:flex; align-items:center; gap:12px; font-size:1.1rem; margin:0 0 8px; }
.num { display:inline-grid; place-items:center; width:32px; height:32px; border-radius:50%;
       background:var(--accent); color:#fff; font-size:.95rem; flex:none; }
.step p { margin:0 0 14px; }
.step img { display:block; max-width:100%; height:auto; border-radius:4px;
            box-shadow:0 1px 4px rgba(0,0,0,.12); }
@media print { main { padding:0; } .step img { box-shadow:none; } }
"""


def _html_inline(text: str) -> str:
    out = []
    for t, bold, italic in inline_parts(text):
        t = html.escape(t)
        out.append(f"<strong>{t}</strong>" if bold else f"<em>{t}</em>" if italic else t)
    return "".join(out)


def export_html(project: Project) -> Path:
    out = export_dir(project)
    title = html.escape(project.display_title())
    body = [f"<h1>{title}</h1>", f'<p class="date">{_date_line(project)}</p>']
    if project.intro:
        body.append(f'<p class="intro">{_html_inline(project.intro)}</p>')
    for n, step, img in _rendered(project):
        parts = [f'<section class="step" id="step-{n}">',
                 f'<h2><span class="num">{n}</span>Step {n}</h2>',
                 f"<p>{_html_inline(Project.instruction(step))}</p>"]
        if img is not None:
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            b64 = base64.b64encode(buf.getvalue()).decode("ascii")
            parts.append(f'<img src="data:image/png;base64,{b64}" alt="Screenshot for step {n}" '
                         f'width="{img.width}" height="{img.height}" loading="lazy">')
        parts.append("</section>")
        body.append("\n".join(parts))

    doc = (f'<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
           f'<meta name="viewport" content="width=device-width, initial-scale=1">\n'
           f"<title>{title}</title>\n<style>{_CSS}</style>\n</head>\n<body>\n<main>\n"
           + "\n".join(body) + "\n</main>\n</body>\n</html>\n")
    path = out / f"{slug(project.display_title())}.html"
    path.write_text(doc, encoding="utf-8")
    return path


# ------------------------------------------------------------------------- Word
def export_docx(project: Project) -> Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Word export needs python-docx: pip install python-docx") from exc

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2))
    usable_w, max_h = 17.0, 19.0

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading(project.display_title(), level=0)
    p = doc.add_paragraph()
    r = p.add_run(_date_line(project))
    r.italic = True
    r.font.color.rgb = RGBColor(0x59, 0x63, 0x6E)
    if project.intro:
        _docx_rich(doc.add_paragraph(), project.intro)

    for n, step, img in _rendered(project):
        h = doc.add_heading(f"Step {n}", level=2)
        h.paragraph_format.keep_with_next = True
        para = doc.add_paragraph()
        _docx_rich(para, Project.instruction(step))
        if img is not None:
            para.paragraph_format.keep_with_next = True
            w_cm = usable_w
            if img.height / img.width * w_cm > max_h:
                w_cm = max_h * img.width / img.height
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            buf.seek(0)
            doc.add_picture(buf, width=Cm(w_cm))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    path = export_dir(project) / f"{slug(project.display_title())}.docx"
    doc.save(path)
    return path


def _docx_rich(paragraph, text: str):
    for t, bold, italic in inline_parts(text):
        run = paragraph.add_run(t)
        run.bold = bold or None
        run.italic = italic or None


EXPORTERS = {"docx": export_docx, "md": export_markdown, "html": export_html}


def export(project: Project, fmt: str) -> Path:
    return EXPORTERS[fmt](project)
